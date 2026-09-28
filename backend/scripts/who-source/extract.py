import openpyxl, json

def load(f):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return rows

def by_day_table(rows):
    return {int(r[0]): {"L": float(r[1]), "M": float(r[2]), "S": float(r[3])} for r in rows}

def by_measure_table(rows):
    out = []
    for r in rows:
        out.append({"x": round(float(r[0]), 1), "L": float(r[1]), "M": float(r[2]), "S": float(r[3])})
    return out

lhfa_boys = by_day_table(load("lhfa-boys.xlsx"))
lhfa_girls = by_day_table(load("lhfa-girls.xlsx"))
wfa_boys = by_day_table(load("wfa-boys.xlsx"))
wfa_girls = by_day_table(load("wfa-girls.xlsx"))

def sample_months(day_table):
    out = {}
    for m in range(0, 60):
        day = round(m * 30.4375)
        if day in day_table:
            out[m] = day_table[day]
        else:
            keys = sorted(day_table.keys())
            nearest = min(keys, key=lambda k: abs(k-day))
            out[m] = day_table[nearest]
    return out

lhfa_boys_m = sample_months(lhfa_boys)
lhfa_girls_m = sample_months(lhfa_girls)
wfa_boys_m = sample_months(wfa_boys)
wfa_girls_m = sample_months(wfa_girls)

print("LHFA boys max day:", max(lhfa_boys.keys()), "count:", len(lhfa_boys))
print("Month 18 LHFA boys:", lhfa_boys_m[18])
print("Month 18 WFA boys:", wfa_boys_m[18])

wfl_boys = by_measure_table(load("wfl-boys.xlsx"))
wfl_girls = by_measure_table(load("wfl-girls.xlsx"))
wfh_boys = by_measure_table(load("wfh-boys.xlsx"))
wfh_girls = by_measure_table(load("wfh-girls.xlsx"))

print("WFL boys range:", wfl_boys[0]['x'], "-", wfl_boys[-1]['x'], "count:", len(wfl_boys))
print("WFH boys range:", wfh_boys[0]['x'], "-", wfh_boys[-1]['x'], "count:", len(wfh_boys))

out = {
    "lhfaMonthly": {"male": lhfa_boys_m, "female": lhfa_girls_m},
    "wfaMonthly": {"male": wfa_boys_m, "female": wfa_girls_m},
    "wflByLength": {"male": wfl_boys, "female": wfl_girls},
    "wfhByHeight": {"male": wfh_boys, "female": wfh_girls},
}

with open("who_lms_extracted.json", "w", encoding="utf-8") as f:
    json.dump(out, f)

import os
print("Written who_lms_extracted.json, size bytes:", os.path.getsize("who_lms_extracted.json"))
